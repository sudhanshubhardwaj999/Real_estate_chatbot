import os
import pytesseract
import cv2
from PyPDF2 import PdfReader
import docx
import openpyxl
from sentence_transformers import SentenceTransformer
import numpy as np

# Update this path to where Tesseract is installed on your system
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def extract_text_from_image(image_path):
    """Extract text from images (JPG, PNG)."""
    try:
        image = cv2.imread(image_path)
        text = pytesseract.image_to_string(image)
        return text
    except Exception as e:
        print(f"Error extracting text from image {image_path}: {e}")
        return ""


def extract_text_from_pdf(pdf_path):
    """Extract text from PDFs."""
    try:
        reader = PdfReader(pdf_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    except Exception as e:
        print(f"Error extracting text from PDF {pdf_path}: {e}")
        return ""


def extract_text_from_docx(docx_path):
    """Extract text from DOCX files."""
    try:
        doc = docx.Document(docx_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    except Exception as e:
        print(f"Error extracting text from DOCX {docx_path}: {e}")
        return ""


def extract_text_from_excel(excel_path):
    """Extract text from Excel files."""
    try:
        workbook = openpyxl.load_workbook(excel_path)
        sheet = workbook.active
        text = ""
        for row in sheet.iter_rows(values_only=True):
            text += " ".join(map(str, row)) + " "
        return text
    except Exception as e:
        print(f"Error extracting text from Excel {excel_path}: {e}")
        return ""


def load_and_preprocess_data(data_dir):
    """Load data from different file formats and preprocess it."""
    all_text = []
    for root, dirs, files in os.walk(data_dir):
        for file in files:
            file_path = os.path.join(root, file)
            if file.endswith((".jpg", ".png")):
                text = extract_text_from_image(file_path)
            elif file.endswith(".pdf"):
                text = extract_text_from_pdf(file_path)
            elif file.endswith(".docx"):
                text = extract_text_from_docx(file_path)
            elif file.endswith(".xlsx"):
                text = extract_text_from_excel(file_path)
            else:
                continue
            if text:
                all_text.append(text)
    return all_text


def embed_data(preprocessed_text):
    """Generate embeddings from preprocessed text."""
    # Ensure preprocessed_text is a list of strings
    if not isinstance(preprocessed_text, list) or not all(
        isinstance(item, str) for item in preprocessed_text
    ):
        raise ValueError("preprocessed_text should be a list of strings")

    model = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")
    embeddings = model.encode(preprocessed_text, convert_to_numpy=True)

    # Ensure embeddings are a 2D array
    if embeddings.ndim != 2:
        raise ValueError("Embeddings should be a 2D array")

    return embeddings
